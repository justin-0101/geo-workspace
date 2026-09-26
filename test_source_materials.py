"""Material ingestion tests: real parsing, real local HTTP fetch, no external network."""
import io
import json
import os
import threading
import unittest
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from unittest.mock import patch

from test_workflow_api import APITests
import source_materials as materials
import workspace_store as store


def minimal_pdf(text='Sample EAM delivery 8 weeks'):
    """A real (tiny) PDF with a text object, so pypdf parsing is actually exercised."""
    payload = text.encode('ascii', 'replace')
    objs = [b'<< /Type /Catalog /Pages 2 0 R >>',
            b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
            b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R '
            b'/Resources << /Font << /F1 5 0 R >> >> >>']
    stream = b'BT /F1 18 Tf 72 700 Td (' + payload + b') Tj ET'
    objs.append(b'<< /Length ' + str(len(stream)).encode() + b' >>\nstream\n' + stream + b'\nendstream')
    objs.append(b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>')
    out = b'%PDF-1.4\n'
    offsets = []
    for index, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += str(index).encode() + b' 0 obj\n' + body + b'\nendobj\n'
    xref = len(out)
    out += b'xref\n0 ' + str(len(objs) + 1).encode() + b'\n0000000000 65535 f \n'
    for offset in offsets:
        out += ('%010d 00000 n \n' % offset).encode()
    out += (b'trailer\n<< /Size ' + str(len(objs) + 1).encode() + b' /Root 1 0 R >>\n'
            b'startxref\n' + str(xref).encode() + b'\n%%EOF')
    return out


def docx_bytes():
    from docx import Document
    document = Document()
    document.add_heading('设备资产管理系统说明', 1)
    document.add_paragraph('本系统支持设备台账、点检、预测性维护与备件管理。')
    document.add_paragraph('已通过 ISO 9001 认证，标准交付周期为 8 周。')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = '型号'
    table.cell(0, 1).text = 'EAM-2024'
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def xlsx_bytes():
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = '报价明细'
    sheet.append(['项目', '金额'])
    sheet.append(['软件许可', 120000])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


FONT_CANDIDATES = (r'C:\Windows\Fonts\msyh.ttc', r'C:\Windows\Fonts\simhei.ttf',
                   r'C:\Windows\Fonts\simsun.ttc')
FONT = next((path for path in FONT_CANDIDATES if os.path.exists(path)), '')


def png_bytes(lines):
    """A real PNG with Chinese text, so OCR is genuinely exercised."""
    from PIL import Image, ImageDraw, ImageFont
    font = ImageFont.truetype(FONT, 34)
    image = Image.new('RGB', (1100, 60 + 70 * len(lines)), 'white')
    draw = ImageDraw.Draw(image)
    for index, line in enumerate(lines):
        draw.text((30, 30 + 70 * index), line, font=font, fill='black')
    buffer = io.BytesIO()
    image.save(buffer, format='PNG')
    return buffer.getvalue()


def scanned_pdf_bytes(lines):
    """A one-page PDF whose only content is an image — no text layer at all."""
    import fitz
    from PIL import Image
    png = png_bytes(lines)
    image = Image.open(io.BytesIO(png))
    document = fitz.open()
    page = document.new_page(width=image.width, height=image.height)
    page.insert_image(fitz.Rect(0, 0, image.width, image.height), stream=png)
    data = document.tobytes()
    document.close()
    return data


class MaterialParserTests(unittest.TestCase):
    def test_text_and_gbk(self):
        text, status, _ = materials.extract_bytes('a.txt', '设备资产管理系统，2024年交付。'.encode('gb18030'))
        self.assertEqual(status, 'ok')
        self.assertIn('2024年交付', text)

    def test_html_drops_scripts_nav_and_footer(self):
        html = ('<html><head><style>a{color:red}</style></head><body>'
                '<nav>首页 产品</nav><h1>EAM 选型</h1><p>要看设备台账与点检。</p>'
                '<script>var x=1</script><footer>Copyright 2024</footer></body></html>')
        text, status, _ = materials.extract_bytes('p.html', html.encode('utf8'))
        self.assertEqual(status, 'ok')
        self.assertIn('EAM 选型', text)
        self.assertNotIn('var x', text)
        self.assertNotIn('Copyright', text)
        self.assertNotIn('首页', text)

    def test_docx_reads_paragraphs_and_tables(self):
        text, status, _ = materials.extract_bytes('d.docx', docx_bytes())
        self.assertEqual(status, 'ok')
        self.assertIn('ISO 9001', text)
        self.assertIn('EAM-2024', text)

    def test_xlsx_reads_sheet(self):
        text, status, _ = materials.extract_bytes('q.xlsx', xlsx_bytes())
        self.assertEqual(status, 'ok')
        self.assertIn('报价明细', text)
        self.assertIn('120000', text)

    def test_pdf_extracts_real_text(self):
        text, status, note = materials.extract_bytes('p.pdf', minimal_pdf())
        self.assertEqual(status, 'ok', note)
        self.assertIn('delivery 8 weeks', text)

    def test_image_is_queued_for_ocr(self):
        text, status, note = materials.extract_bytes('x.png', b'\x89PNG\r\n\x1a\n' + b'0' * 64)
        self.assertEqual(text, '')
        self.assertEqual(status, 'ocr_pending')
        self.assertIn('OCR', note)

    def test_image_without_ocr_falls_back_to_unsupported(self):
        with patch.dict(os.environ, {'GEO_OCR': '0'}):
            self.assertFalse(materials.ocr_available())
            text, status, note = materials.extract_bytes('x.png', b'\x89PNG\r\n\x1a\n' + b'0' * 64)
            self.assertEqual((text, status), ('', 'unsupported'))
            self.assertIn('图片', note)

    def test_scanned_pdf_without_ocr_reports_why(self):
        data = scanned_pdf_bytes(['标准交付周期为 8 周'])
        self.assertEqual(materials.extract_bytes('scan.pdf', data)[1], 'ocr_pending')
        with patch.dict(os.environ, {'GEO_OCR': '0'}):
            text, status, note = materials.extract_bytes('scan.pdf', data)
            self.assertEqual((text, status), ('', 'empty'))
            self.assertIn('没有文字层', note)
            self.assertIn('OCR', note)

    def test_broken_file_fails_without_raising(self):
        for name, data in (('bad.pdf', b'not a pdf'), ('bad.docx', b'not a docx'),
                           ('bad.xlsx', b'not a sheet')):
            text, status, note = materials.extract_bytes(name, data)
            self.assertEqual(text, '', name)
            self.assertIn(status, {'failed', 'empty'}, name)
            self.assertTrue(note, name)

    def test_empty_file_reports_empty(self):
        self.assertEqual(materials.extract_bytes('e.txt', b'')[1], 'empty')

    def test_unknown_extension_falls_back_to_text(self):
        text, status, note = materials.extract_bytes('a.bin', 'plain 2024 text'.encode())
        self.assertEqual(status, 'ok')
        self.assertIn('未识别的扩展名', note)


class _Handler(BaseHTTPRequestHandler):
    payloads = {}

    def do_GET(self):
        item = self.payloads.get(self.path)
        if not item:
            self.send_response(404)
            self.end_headers()
            return
        body, ctype = item
        self.send_response(200)
        self.send_header('Content-Type', ctype)
        self.send_header('Content-Length', str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, *args):
        pass


class UrlFetchTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        _Handler.payloads = {
            '/about': ('<html><nav>首页</nav><body><h1>关于我们</h1>'
                       '<p>公司成立于2003年，现有员工 260 人。</p>'
                       '<script>track()</script></body></html>'.encode('utf8'), 'text/html; charset=utf-8'),
            '/plain.txt': ('交付周期为 8 周。'.encode('utf8'), 'text/plain; charset=utf-8'),
            '/report.pdf': (minimal_pdf('Annual report 2024 revenue'), 'application/pdf'),
            '/blank': (b'<html><body></body></html>', 'text/html'),
            '/binary.bin': (b'\x00\x01\x02\x03', 'application/octet-stream'),
        }
        cls.server = ThreadingHTTPServer(('127.0.0.1', 0), _Handler)
        cls.port = cls.server.server_address[1]
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()

    def setUp(self):
        # The loopback guard is a deliberate product default; tests opt out explicitly.
        self.env = patch.dict(os.environ, {'GEO_ALLOW_LOCAL_FETCH': '1'})
        self.env.start()

    def tearDown(self):
        self.env.stop()

    def url(self, path):
        return f'http://127.0.0.1:{self.port}{path}'

    def test_html_page_is_extracted(self):
        result = materials.fetch_url(self.url('/about'))
        self.assertEqual(result['status'], 'ok')
        self.assertIn('公司成立于2003年', result['text'])
        self.assertNotIn('track()', result['text'])
        self.assertNotIn('首页', result['text'])

    def test_text_plain_is_extracted(self):
        self.assertIn('8 周', materials.fetch_url(self.url('/plain.txt'))['text'])

    def test_pdf_url_is_extracted(self):
        result = materials.fetch_url(self.url('/report.pdf'))
        self.assertEqual(result['status'], 'ok')
        self.assertIn('revenue', result['text'])

    def test_empty_page_reports_empty_with_reason(self):
        result = materials.fetch_url(self.url('/blank'))
        self.assertEqual(result['status'], 'empty')
        self.assertIn('没有可提取的文字', result['note'])

    def test_unsupported_content_type_is_reported(self):
        result = materials.fetch_url(self.url('/binary.bin'))
        self.assertEqual(result['status'], 'unsupported')
        self.assertIn('application/octet-stream', result['note'])

    def test_404_becomes_readable_error(self):
        with self.assertRaises(materials.MaterialError) as ctx:
            materials.fetch_url(self.url('/missing'))
        self.assertIn('404', str(ctx.exception))

    def test_loopback_is_blocked_by_default(self):
        with patch.dict(os.environ, {'GEO_ALLOW_LOCAL_FETCH': ''}):
            with self.assertRaises(materials.MaterialError) as ctx:
                materials.fetch_url('http://127.0.0.1:8080/x')
            self.assertIn('本机或内网', str(ctx.exception))

    def test_bad_scheme_is_rejected(self):
        for bad in ('file:///etc/passwd', 'javascript:alert(1)', 'notaurl', ''):
            with self.assertRaises(materials.MaterialError, msg=bad):
                materials.fetch_url(bad)


class MaterialApiTests(APITests):
    def tearDown(self):
        # 后台 OCR 线程必须在临时库被换掉之前结束。
        materials.wait_for_ocr()
        super().tearDown()

    def asset(self):
        slug = self.create('素材项目')
        base = '/api/projects/' + slug
        created = self.client.post(base + '/content', json={'title': '素材测试内容'})
        return slug, base, created.json()['id']

    def test_upload_note_list_refresh_delete_download(self):
        slug, base, aid = self.asset()
        url = base + '/content/' + aid + '/sources'

        upload = self.client.post(url + '/file',
                                  files=[('files', ('说明.md', '# 说明\n本系统支持设备台账与点检。'.encode('utf8'),
                                                    'text/markdown'))])
        self.assertEqual(upload.status_code, 200, upload.text)
        result = upload.json()['results'][0]
        self.assertEqual(result['status'], 'ok')
        self.assertGreater(result['chars'], 5)

        note = self.client.post(url + '/note', json={'text': '交付周期为 8 周。', 'label': '销售口述'})
        self.assertEqual(note.status_code, 201)

        items = self.client.get(url).json()['items']
        self.assertEqual(len(items), 2)
        self.assertEqual([x['kind'] for x in items], ['file', 'note'])
        self.assertNotIn('extracted_text', items[0])          # 列表不回传全文
        self.assertGreater(items[0]['chars'], 0)

        download = self.client.get(f"{url}/{result['id']}/download")
        self.assertEqual(download.status_code, 200)
        self.assertIn('设备台账', download.content.decode('utf8'))

        refreshed = self.client.post(f"{url}/{result['id']}/refresh")
        self.assertEqual(refreshed.status_code, 200)
        self.assertEqual(refreshed.json()['status'], 'ok')

        self.assertEqual(self.client.delete(f"{url}/{result['id']}").status_code, 200)
        self.assertEqual(len(self.client.get(url).json()['items']), 1)

    def test_upload_rejects_empty_file_without_killing_the_batch(self):
        slug, base, aid = self.asset()
        url = base + '/content/' + aid + '/sources/file'
        out = self.client.post(url, files=[
            ('files', ('good.md', '有效内容 2024 年交付。'.encode('utf8'), 'text/markdown')),
            ('files', ('empty.md', b'', 'text/markdown')),
        ])
        self.assertEqual(out.status_code, 200)
        statuses = [(r['file_name'], r.get('status') or 'failed') for r in out.json()['results']]
        self.assertEqual(statuses, [('good.md', 'ok'), ('empty.md', 'failed')])
        self.assertEqual(len(self.client.get(base + '/content/' + aid + '/sources').json()['items']), 1)

    def test_url_material_is_extracted_into_the_asset(self):
        slug, base, aid = self.asset()
        url = base + '/content/' + aid + '/sources/url'
        with patch.dict(os.environ, {'GEO_ALLOW_LOCAL_FETCH': '1'}):
            server = ThreadingHTTPServer(('127.0.0.1', 0), _Handler)
            port = server.server_address[1]
            threading.Thread(target=server.serve_forever, daemon=True).start()
            _Handler.payloads = {'/about': ('<html><body><p>公司成立于2003年，员工 260 人。</p></body></html>'
                                            .encode('utf8'), 'text/html; charset=utf-8')}
            try:
                out = self.client.post(url, json={'url': f'http://127.0.0.1:{port}/about', 'label': '官网关于页'})
                self.assertEqual(out.status_code, 201, out.text)
                self.assertEqual(out.json()['status'], 'ok')
            finally:
                server.shutdown()
                server.server_close()

        context = self.client.get(base + '/content/' + aid + '/context').json()
        self.assertEqual(len(context['sources']), 1)
        bundle = context['source_bundle']
        self.assertEqual(len(bundle['materials']), 1)
        self.assertIn('公司成立于2003年', bundle['materials_text'])

    def test_upload_preflight_is_allowed_from_the_local_frontend(self):
        """前端在 4173、API 在 8798，属于跨源；multipart 上传会先发预检。"""
        slug, base, aid = self.asset()
        url = base + '/content/' + aid + '/sources/file'
        response = self.client.options(url, headers={
            'Origin': 'http://127.0.0.1:4173',
            'Access-Control-Request-Method': 'POST',
            'Access-Control-Request-Headers': 'content-type'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers.get('access-control-allow-origin'), 'http://127.0.0.1:4173')
        self.assertIn('content-type', (response.headers.get('access-control-allow-headers') or '').lower())

    def test_sources_are_isolated_between_projects(self):
        slug, base, aid = self.asset()
        self.client.post(base + '/content/' + aid + '/sources/note', json={'text': '隔离测试内容 2024。'})
        other = self.create('另一个项目')
        other_base = '/api/projects/' + other
        other_asset = self.client.post(other_base + '/content', json={'title': '别的项目'}).json()['id']
        self.assertEqual(self.client.get(other_base + '/content/' + other_asset + '/sources').json()['items'], [])
        # 跨项目不能读到、也不能删到别人的素材
        source = self.client.get(base + '/content/' + aid + '/sources').json()['items'][0]
        self.assertEqual(self.client.delete(
            other_base + '/content/' + other_asset + '/sources/' + source['id']).status_code, 404)
        self.assertEqual(self.client.post(
            other_base + '/content/' + other_asset + '/sources/' + source['id'] + '/refresh').status_code, 404)

    def test_deleting_asset_cascades_to_sources(self):
        slug, base, aid = self.asset()
        self.client.post(base + '/content/' + aid + '/sources/note', json={'text': '内容 2024。'})
        with store.connection() as c:
            c.execute('DELETE FROM editorial_assets WHERE id=?', (aid,))
        self.assertEqual(materials.list_sources(slug, aid), [])

    def test_generation_uses_materials_and_fills_facts(self):
        slug, base, aid = self.asset()
        self.client.post(base + '/content/' + aid + '/sources/note',
                         json={'text': '本系统支持设备台账、点检、预测性维护与备件管理。\n'
                                       '已通过 ISO 9001 认证，标准交付周期为 8 周。\n'
                                       '实施配置包含数据治理与接口集成，交付后提供用户培训。\n'
                                       '设备台账反映设备当前状态，支持标签打印与批量导入。\n'
                                       '点检与保养计划可按设备类别下发，并保留执行记录。\n'
                                       '备件管理支持安全库存预警与出入库台账。',
                               'label': '产品说明'})
        out = self.client.post(base + '/content/' + aid + '/generate',
                              json={'revision': 1, 'channel': '公众号长文', 'audience': '制造企业',
                                    'objective': '帮助选型'})
        self.assertEqual(out.status_code, 200, out.text)
        asset = next(x for x in self.client.get(base + '/editorial/content').json()['items'] if x['id'] == aid)
        self.assertIn('ISO 9001 认证', asset['body'])          # 素材内容进了正文
        self.assertIn('粘贴素材「产品说明」', asset['facts'])   # 来源回填到事实依据
        self.assertIn('待人工确认', asset['facts'])
        meta = json.loads(asset['generation_meta_json'])
        self.assertEqual(meta['material_chars'] > 0, True)
        self.assertEqual(len(meta['material_ids']), 1)
        self.assertTrue(out.json()['quality']['passed'], out.json()['quality']['errors'])

    def test_materials_are_not_silently_promoted_to_verified_facts(self):
        """素材只是写作输入：它不能自己变成“已确认事实”，也不能绕过审核。"""
        slug, base, aid = self.asset()
        self.client.post(base + '/content/' + aid + '/sources/note', json={'text': '公司成立于2003年。'})
        context = self.client.get(base + '/content/' + aid + '/context').json()
        self.assertEqual(context['source_bundle']['verified_facts'], '')
        self.assertIn('仍需人工确认', context['source_bundle']['material_notice'])
        self.assertFalse(context['quality']['passed'])          # 还没生成正文，不能过审

    def test_library_lists_every_item_with_index_and_metadata(self):
        slug = self.create('内容库项目')
        base = '/api/projects/' + slug
        first = self.client.post(base + '/content', json={'title': '第一篇', 'channel': '公众号长文'}).json()['id']
        second = self.client.post(base + '/content', json={'title': '第二篇', 'channel': '知乎回答'}).json()['id']
        self.client.post(base + '/content/' + first + '/sources/note', json={'text': '公司成立于2003年，员工 260 人。'})
        self.client.post(base + '/content/' + first + '/generate', json={'revision': 1, 'channel': '公众号长文'})

        items = self.client.get(base + '/library').json()['items']
        self.assertEqual([x['title'] for x in items], ['第二篇', '第一篇'])   # 最新在前
        self.assertEqual([x['index'] for x in items], [1, 2])                # 序号按列表顺序
        newest, oldest = items
        self.assertEqual(newest['generated_at'], '')                          # 没生成过
        self.assertEqual(newest['material_count'], 0)
        self.assertTrue(oldest['generated_at'])                               # 生成过
        self.assertEqual(oldest['engine'], 'local-safe')
        self.assertEqual(oldest['material_count'], 1)
        self.assertGreater(oldest['chars'], 100)
        self.assertNotIn('body', oldest)                                      # 列表不带正文

    def test_delete_content_removes_materials_and_frees_the_action(self):
        slug, base, aid = self.asset()
        self.client.post(base + '/content/' + aid + '/sources/note', json={'text': '会被一起删掉的素材。'})
        removed = self.client.delete(base + '/content/' + aid)
        self.assertEqual(removed.status_code, 200, removed.text)
        self.assertEqual(removed.json()['deleted_sources'], 1)
        self.assertEqual(self.client.get(base + '/library').json()['items'], [])
        self.assertEqual(self.client.get(base + '/content/' + aid + '/context').status_code, 404)

    def test_delete_content_drops_material_files_from_disk(self):
        slug, base, aid = self.asset()
        upload = self.client.post(base + '/content/' + aid + '/sources/file',
                                  files=[('files', ('a.md', '设备台账与点检。'.encode('utf8'), 'text/markdown'))])
        self.assertEqual(upload.status_code, 200)
        with store.connection() as c:
            row = dict(c.execute('SELECT * FROM asset_sources').fetchone())
        path = store.DATA / row['file_path']
        self.assertTrue(path.is_file())
        self.client.delete(base + '/content/' + aid)
        self.assertFalse(path.is_file())

    def test_partial_save_does_not_wipe_auto_filled_facts(self):
        """界面不再手填事实依据；编辑正文不能把自动带出的来源清空。"""
        slug, base, aid = self.asset()
        self.client.post(base + '/content/' + aid + '/sources/note', json={'text': '公司成立于2003年，员工 260 人。'})
        self.client.post(base + '/content/' + aid + '/generate', json={'revision': 1, 'channel': '公众号长文'})
        before = self.client.get(base + '/editorial/content').json()['items'][0]
        self.assertIn('粘贴素材', before['facts'])
        saved = self.client.put(base + '/content/' + aid,
                                json={'title': '改过的标题', 'body': before['body'], 'revision': before['revision']})
        self.assertEqual(saved.status_code, 200, saved.text)
        after = self.client.get(base + '/editorial/content').json()['items'][0]
        self.assertEqual(after['facts'], before['facts'])          # 来源保留
        self.assertEqual(after['summary'], before['summary'])      # 摘要也没被清空


@unittest.skipUnless(materials.ocr_available() and FONT, '本机没有可用的 RapidOCR 或中文字体，跳过 OCR 用例')
class OcrTests(MaterialApiTests):
    """真图 + 真扫描件的 OCR 验证（本机识别，不联网）。"""

    def upload_and_wait(self, slug, aid, name, data):
        response = self.client.post('/api/projects/' + slug + '/content/' + aid + '/sources/file',
                                    files=[('files', (name, data, 'application/octet-stream'))])
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()['results'][0]['status'], 'ocr_pending')
        self.assertTrue(materials.wait_for_ocr(), 'OCR 没有在超时内结束')
        items = self.client.get('/api/projects/' + slug + '/content/' + aid + '/sources').json()['items']
        self.assertEqual(len(items), 1)
        return items[0]

    def test_image_upload_is_recognised(self):
        slug, base, aid = self.asset()
        item = self.upload_and_wait(slug, aid, '扫描件.png',
                                    png_bytes(['示例科技有限公司',
                                               '已通过 ISO 9001 认证']))
        self.assertEqual(item['status'], 'ok', item['note'])
        self.assertGreater(item['chars'], 5)
        self.assertIn('OCR', item['note'])
        bundle = self.client.get(base + '/content/' + aid + '/context').json()['source_bundle']
        self.assertIn('ISO', bundle['materials_text'])          # 识别结果真的进了生成输入
        self.assertIn('示例科技', bundle['materials_text'])

    def test_scanned_pdf_is_recognised(self):
        slug, base, aid = self.asset()
        item = self.upload_and_wait(slug, aid, '扫描报价单.pdf',
                                    scanned_pdf_bytes(['标准交付周期为 8 周',
                                                       '实施配置包含数据治理']))
        self.assertEqual(item['status'], 'ok', item['note'])
        bundle = self.client.get(base + '/content/' + aid + '/context').json()['source_bundle']
        self.assertIn('数据治理', bundle['materials_text'])

    def test_ocr_text_reaches_the_generated_draft(self):
        slug, base, aid = self.asset()
        self.upload_and_wait(slug, aid, '扫描件.png', png_bytes(['已通过 ISO 9001 认证']))
        out = self.client.post(base + '/content/' + aid + '/generate',
                               json={'revision': 1, 'channel': '公众号长文'})
        self.assertEqual(out.status_code, 200, out.text)
        asset = self.client.get(base + '/editorial/content').json()['items'][0]
        # OCR 的空格不稳定（实测输出“已通过ISO9001认证”），比较前先去空格。
        self.assertIn('ISO9001', asset['body'].replace(' ', ''))
        self.assertIn('扫描件.png', asset['facts'])

    def test_pending_ocr_is_reported_as_a_generation_warning(self):
        """素材还在识别时生成，必须明确告知本次没用到它，不能静默忽略。"""
        slug, base, aid = self.asset()
        with patch.object(materials, 'queue_ocr', lambda *a, **k: False):
            self.client.post(base + '/content/' + aid + '/sources/file',
                             files=[('files', ('慢.png', png_bytes(['ISO 9001']), 'image/png'))])
            out = self.client.post(base + '/content/' + aid + '/generate',
                                   json={'revision': 1, 'channel': '公众号长文'})
        self.assertEqual(out.status_code, 200, out.text)
        self.assertTrue(any('OCR' in w for w in out.json()['warnings']), out.json()['warnings'])

    def test_restart_resets_stuck_ocr_rows(self):
        slug, base, aid = self.asset()
        with patch.object(materials, 'queue_ocr', lambda *a, **k: False):
            self.client.post(base + '/content/' + aid + '/sources/file',
                             files=[('files', ('卡住.png', png_bytes(['ISO 9001']), 'image/png'))])
        with store.connection() as c:
            c.execute("UPDATE asset_sources SET extract_status='ocr_pending'")
        self.assertEqual(materials.reset_stuck_ocr(), 1)
        item = self.client.get(base + '/content/' + aid + '/sources').json()['items'][0]
        self.assertEqual(item['status'], 'pending')
        self.assertIn('重新提取', item['note'])

    def test_refresh_requeues_ocr(self):
        slug, base, aid = self.asset()
        item = self.upload_and_wait(slug, aid, '扫描件.png', png_bytes(['设备资产管理系统']))
        refreshed = self.client.post(f"{base}/content/{aid}/sources/{item['id']}/refresh")
        self.assertEqual(refreshed.status_code, 200)
        self.assertEqual(refreshed.json()['status'], 'ocr_pending')
        self.assertTrue(materials.wait_for_ocr())
        again = self.client.get(base + '/content/' + aid + '/sources').json()['items'][0]
        self.assertEqual(again['status'], 'ok')


if __name__ == '__main__':
    unittest.main()
