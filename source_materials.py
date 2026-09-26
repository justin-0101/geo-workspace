"""Content source materials: uploads, URLs, pasted text and OCR.

Extracted text is raw drafting input. It is never treated as a confirmed business fact —
the editor still confirms facts and sources before a version can be reviewed.

OCR (RapidOCR + ONNX, purely local) runs on a background thread: recognition costs
roughly 8 seconds per page on this machine, so uploads must not block on it. A material
moves through ``ocr_pending`` → ``ok`` and the UI polls until it settles.

Parsers are imported lazily: a missing optional library must degrade into a readable
"无法解析" note on one material, never break the whole content module.
"""
import io
import os
import re
import threading
import time
import uuid
from pathlib import Path
from urllib import error, request
from urllib.parse import urlparse

import workspace_store as store

KINDS = ('file', 'url', 'note')
STATUSES = ('ok', 'empty', 'failed', 'unsupported', 'ocr_pending')

MAX_UPLOAD_BYTES = 20 * 1024 * 1024
MAX_URL_BYTES = 3 * 1024 * 1024
MAX_STORED_CHARS = 240_000
USER_AGENT = ('Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 '
              '(KHTML, like Gecko) Chrome/124.0 Safari/537.36')

TEXT_SUFFIXES = {'.txt', '.md', '.markdown', '.csv', '.tsv', '.json', '.xml', '.log', '.srt'}
HTML_SUFFIXES = {'.html', '.htm', '.xhtml'}
IMAGE_SUFFIXES = {'.png', '.jpg', '.jpeg', '.webp', '.gif', '.bmp', '.tif', '.tiff'}
OFFICE_SUFFIXES = {'.pdf', '.docx', '.xlsx', '.xlsm'}

#: 只拦回环与链路本地地址，避免误把本机服务当成外部素材抓。
#: 这是字符串判断，不做 DNS 解析；域名解析到 127.0.0.1 的情况拦不住。
#: 需要故意抓本机/内网服务（或跑集成测试）时，设 GEO_ALLOW_LOCAL_FETCH=1 放开。
BLOCKED_HOST_PREFIXES = ('localhost', '127.', '0.0.0.0', '::1', '169.254.')
BLOCKED_HOST_SUFFIXES = ('.local', '.localhost', '.internal')


def local_fetch_allowed():
    return str(os.environ.get('GEO_ALLOW_LOCAL_FETCH') or '').strip().lower() in {'1', 'true', 'yes'}


class MaterialError(ValueError):
    """用户可读的素材错误。"""


def _env_int(name, default, low, high):
    try:
        value = int(str(os.environ.get(name) or '').strip())
    except (TypeError, ValueError):
        return default
    return max(low, min(value, high))


def ocr_max_pages():
    return _env_int('GEO_OCR_MAX_PAGES', 6, 1, 80)


def ocr_dpi():
    return _env_int('GEO_OCR_DPI', 150, 72, 400)


def ocr_enabled():
    return str(os.environ.get('GEO_OCR') or '1').strip().lower() not in {'0', 'false', 'no'}


# --------------------------------------------------------------------------- 文本处理

def _normalize(text):
    text = str(text or '').replace('\r\n', '\n').replace('\r', '\n')
    text = re.sub(r'[ \t\x0b\f\xa0\u3000]+', ' ', text)
    lines = [line.strip() for line in text.split('\n')]
    out = []
    for line in lines:
        if line or (out and out[-1]):
            out.append(line)
    return '\n'.join(out).strip()


def decode_bytes(data):
    for encoding in ('utf-8-sig', 'utf-8', 'gb18030', 'big5'):
        try:
            return data.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode('latin-1', errors='replace')


def html_to_text(data):
    raw = decode_bytes(data) if isinstance(data, bytes) else str(data or '')
    try:
        from lxml import html as lxml_html
        doc = lxml_html.fromstring(raw)
        for bad in doc.xpath('//script|//style|//noscript|//svg|//template|//nav|//footer|//aside|//form|//iframe'):
            parent = bad.getparent()
            if parent is not None:
                parent.remove(bad)
        text = doc.text_content()
    except Exception:
        # 解析器不可用或页面畸形时退化为去标签，宁可粗糙也不要整条素材失败。
        text = re.sub(r'(?is)<(script|style)[^>]*>.*?</\1>', ' ', raw)
        text = re.sub(r'(?s)<[^>]+>', ' ', text)
        text = (text.replace('&nbsp;', ' ').replace('&amp;', '&')
                    .replace('&lt;', '<').replace('&gt;', '>').replace('&quot;', '"'))
    return _normalize(text)


def _truncate(text, note=''):
    if len(text) > MAX_STORED_CHARS:
        return text[:MAX_STORED_CHARS], (note + f'；文本过长，只保留前 {MAX_STORED_CHARS} 字').strip('；')
    return text, note


# --------------------------------------------------------------------------- OCR

#: RapidOCR 的模型加载与推理都吃满 CPU，同一时刻只跑一个，避免互相拖慢。
_OCR_LOCK = threading.RLock()
_OCR_ENGINE = None
_OCR_JOBS = set()
_OCR_THREADS = {}
_OCR_JOBS_LOCK = threading.Lock()


def ocr_available():
    if not ocr_enabled():
        return False
    try:
        import rapidocr_onnxruntime  # noqa: F401
        return True
    except Exception:
        return False


def _engine():
    global _OCR_ENGINE
    if _OCR_ENGINE is None:
        from rapidocr_onnxruntime import RapidOCR
        _OCR_ENGINE = RapidOCR()
    return _OCR_ENGINE


def _run_engine(images):
    """images: list of BGR numpy arrays. Returns joined text."""
    with _OCR_LOCK:
        engine = _engine()
        lines = []
        for image in images:
            output = engine(image)
            result = output[0] if isinstance(output, tuple) else output
            for item in (result or []):
                if isinstance(item, (list, tuple)) and len(item) > 1 and item[1]:
                    lines.append(str(item[1]).strip())
    return _normalize('\n'.join(line for line in lines if line))


def _image_to_array(data):
    import cv2
    import numpy as np
    array = cv2.imdecode(np.frombuffer(data, np.uint8), cv2.IMREAD_COLOR)
    if array is None:
        raise MaterialError('图片无法解码（可能不是真正的图片文件）')
    return array


def _pdf_page_arrays(data, limit, dpi):
    import fitz
    import numpy as np
    document = fitz.open(stream=data, filetype='pdf')
    try:
        total = document.page_count
        arrays = []
        for index in range(min(total, limit)):
            pixmap = document[index].get_pixmap(dpi=dpi)
            array = np.frombuffer(pixmap.samples, dtype=np.uint8).reshape(
                pixmap.height, pixmap.width, pixmap.n)
            if pixmap.n == 4:
                array = array[:, :, :3]
            arrays.append(array[:, :, ::-1].copy())      # RGB → BGR
        return arrays, total
    finally:
        document.close()


def ocr_bytes(filename, data, limit=None, dpi=None):
    """OCR an image or a PDF. Returns (text, status, note)."""
    limit = limit or ocr_max_pages()
    dpi = dpi or ocr_dpi()
    suffix = Path(str(filename or '')).suffix.lower()
    try:
        if suffix == '.pdf':
            arrays, total = _pdf_page_arrays(data, limit, dpi)
            scope = f'共 {total} 页' + ('' if total <= limit else f'，只识别前 {limit} 页')
        else:
            arrays = [_image_to_array(data)]
            scope = ''
    except MaterialError:
        raise
    except ImportError as exc:
        return '', 'failed', f'缺少 OCR 依赖 {exc.name}'
    except Exception as exc:
        return '', 'failed', f'OCR 准备失败：{type(exc).__name__}: {exc}'[:300]
    if not arrays:
        return '', 'empty', '没有可识别的页面'
    try:
        text = _run_engine(arrays)
    except ImportError as exc:
        return '', 'failed', f'缺少 OCR 依赖 {exc.name}'
    except Exception as exc:
        return '', 'failed', f'OCR 失败：{type(exc).__name__}: {exc}'[:300]
    note = '·'.join(x for x in (scope, f'OCR 识别 {len(arrays)} 页，可能有错字') if x)
    text, cut = _truncate(text)
    note = '；'.join(x for x in (note, cut) if x)
    if not text:
        return '', 'empty', note + '；OCR 没有识别出文字'
    return text, 'ok', note


# --------------------------------------------------------------------------- 文件解析

def _extract_pdf(data):
    """Returns (text, status, note, needs_ocr)."""
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    if getattr(reader, 'is_encrypted', False):
        try:
            reader.decrypt('')
        except Exception:
            return '', 'failed', 'PDF 有密码，无法提取文字', False
    total = len(reader.pages)
    limit = min(total, 80)
    chunks = []
    for index in range(limit):
        try:
            chunks.append(reader.pages[index].extract_text() or '')
        except Exception as exc:
            chunks.append('')
            if index == 0:
                return '', 'failed', f'PDF 解析失败：{type(exc).__name__}', False
    text = _normalize('\n'.join(chunks))
    note = f'共 {total} 页' + ('' if limit == total else f'，只提取前 {limit} 页')
    if text:
        return text, 'ok', note, False
    # 没有文字层：交给 OCR（可能是扫描件）
    return '', 'empty', note + '；没有文字层（可能是扫描件）', True


def _extract_docx(data):
    from docx import Document
    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(' | '.join(cells))
    text = _normalize('\n'.join(parts))
    if not text:
        return '', 'empty', '文档里没有可提取的文字'
    return text, 'ok', ''


def _extract_xlsx(data):
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for sheet in list(book.worksheets)[:8]:
        parts.append(f'# 工作表 {sheet.title}')
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            if index >= 400:
                parts.append('…（行数过多，已截断）')
                break
            cells = [str(cell).strip() for cell in row if cell not in (None, '')]
            if cells:
                parts.append(' | '.join(cells))
    text = _normalize('\n'.join(parts))
    if not text:
        return '', 'empty', '表格里没有可提取的内容'
    return text, 'ok', ''


def extract_bytes(filename, data):
    """Parse immediately available text.

    Returns (text, status, note) where status may be ``ocr_pending`` meaning the caller
    should hand the bytes to :func:`run_ocr`.
    """
    name = str(filename or '')
    suffix = Path(name).suffix.lower()
    if isinstance(data, str):
        data = data.encode('utf-8')
    try:
        if suffix == '.pdf':
            text, status, note, needs_ocr = _extract_pdf(data)
            if needs_ocr:
                if ocr_available():
                    return '', 'ocr_pending', note + f'；已排队 OCR（最多 {ocr_max_pages()} 页）'
                return '', 'empty', note + '；本机没有可用的 OCR，请人工录入或改用可复制文字的版本'
            if text:
                text, cut = _truncate(text)
                return text, status, '；'.join(x for x in (note, cut) if x)
            return text, status, note
        if suffix == '.docx':
            return _extract_docx(data)
        if suffix in {'.xlsx', '.xlsm'}:
            return _extract_xlsx(data)
        if suffix in HTML_SUFFIXES:
            text, note = _truncate(html_to_text(data))
            return (text, 'ok', note) if text else ('', 'empty', '页面没有可提取的文字')
        if suffix in IMAGE_SUFFIXES:
            if ocr_available():
                return '', 'ocr_pending', '图片已排队 OCR 识别'
            return '', 'unsupported', '图片无法自动提取文字（本机没有可用的 OCR），请在名称里写清它要说明的事实'
        text, note = _truncate(_normalize(decode_bytes(data)))
        if suffix not in TEXT_SUFFIXES and suffix:
            note = (note + f'；按纯文本读取（未识别的扩展名 {suffix}）').strip('；')
        return (text, 'ok', note) if text else ('', 'empty', '文件没有可提取的文字')
    except ImportError as exc:
        return '', 'failed', f'缺少解析库 {exc.name}，无法解析 {suffix or "该文件"}'
    except Exception as exc:
        return '', 'failed', f'解析失败：{type(exc).__name__}: {exc}'[:300]


# --------------------------------------------------------------------------- 网址抓取

def _check_host(hostname):
    host = (hostname or '').lower()
    if not host:
        raise MaterialError('网址缺少域名')
    if local_fetch_allowed():
        return
    if host.startswith(BLOCKED_HOST_PREFIXES) or host.endswith(BLOCKED_HOST_SUFFIXES):
        raise MaterialError('不支持抓取本机或内网地址（如确需，请设 GEO_ALLOW_LOCAL_FETCH=1）；也可改用上传附件或粘贴文本')


def _opener():
    return request.build_opener(request.ProxyHandler({}))


def fetch_url(url):
    """Return dict(text, status, note, final_url, content_type, raw, suffix)."""
    parsed = urlparse(str(url or '').strip())
    if parsed.scheme not in {'http', 'https'} or not parsed.hostname:
        raise MaterialError('请填写完整的 http/https 网址')
    _check_host(parsed.hostname)
    req = request.Request(parsed.geturl(), headers={
        'User-Agent': USER_AGENT,
        'Accept': 'text/html,application/xhtml+xml,application/pdf;q=0.9,image/*;q=0.8,*/*;q=0.8',
        'Accept-Language': 'zh-CN,zh;q=0.9,en;q=0.8',
    })
    try:
        with _opener().open(req, timeout=25) as response:
            ctype = (response.headers.get('Content-Type') or '').split(';')[0].strip().lower()
            raw = response.read(MAX_URL_BYTES + 1)
            final_url = response.geturl()
    except error.HTTPError as exc:
        raise MaterialError(f'抓取失败：HTTP {exc.code}（{exc.reason}）') from exc
    except error.URLError as exc:
        raise MaterialError(f'抓取失败：{exc.reason}') from exc
    except Exception as exc:
        raise MaterialError(f'抓取失败：{type(exc).__name__}: {exc}'[:300]) from exc

    truncated = len(raw) > MAX_URL_BYTES
    raw = raw[:MAX_URL_BYTES]
    note = '页面较大，只抓取了前 3MB' if truncated else ''
    suffix = Path(urlparse(final_url).path).suffix.lower()
    text = ''
    if ctype == 'application/pdf' or suffix == '.pdf':
        text, status, extra, needs_ocr = _extract_pdf(raw)
        if needs_ocr:
            status = 'ocr_pending' if ocr_available() else 'empty'
            extra = extra + (f'；已排队 OCR（最多 {ocr_max_pages()} 页）' if ocr_available()
                             else '；本机没有可用的 OCR')
    elif ctype in {'text/html', 'application/xhtml+xml'} or suffix in HTML_SUFFIXES:
        text = html_to_text(raw)
        status, extra = ('ok', '') if text else ('empty', '页面没有可提取的文字（可能是需要登录或纯前端渲染的页面）')
    elif ctype.startswith('image/') or suffix in IMAGE_SUFFIXES:
        status = 'ocr_pending' if ocr_available() else 'unsupported'
        extra = '图片已排队 OCR 识别' if ocr_available() else '图片无法自动提取文字（本机没有可用的 OCR）'
    elif ctype.startswith('text/') or ctype in {'application/json', 'application/xml', 'application/rss+xml'}:
        text = _normalize(decode_bytes(raw))
        status, extra = ('ok', '') if text else ('empty', '响应没有可提取的文字')
    else:
        status, extra = 'unsupported', f'不支持的响应类型 {ctype or "未知"}'
    text, cut = _truncate(text)
    return {'text': text, 'status': status,
            'note': '；'.join(x for x in (note, cut, extra) if x)[:400],
            'final_url': final_url, 'content_type': ctype, 'raw': raw, 'suffix': suffix}


# --------------------------------------------------------------------------- 落盘与持久化

def storage_dir(slug, asset_id):
    return store.DATA / 'content' / slug / asset_id


def _safe_suffix(filename):
    suffix = Path(str(filename or '')).suffix.lower()
    return suffix if re.fullmatch(r'\.[a-z0-9]{1,8}', suffix or '') else ''


def _safe_label(name, fallback='素材'):
    cleaned = re.sub(r'[\\/:*?"<>|\x00-\x1f]+', '_', str(name or '')).strip(' .')
    return cleaned[:120] or fallback


def _write_file(slug, asset_id, source_id, filename, data):
    directory = storage_dir(slug, asset_id)
    directory.mkdir(parents=True, exist_ok=True)
    target = directory / (source_id + _safe_suffix(filename))
    resolved = target.resolve()
    if store.DATA.resolve() not in resolved.parents:
        raise MaterialError('素材路径不正确')
    target.write_bytes(data)
    return str(target.relative_to(store.DATA)).replace('\\', '/')


def _asset_exists(c, slug, asset_id):
    if not c.execute('SELECT 1 FROM editorial_assets WHERE id=? AND project_slug=?',
                     (asset_id, slug)).fetchone():
        raise KeyError('内容不存在')


def _insert(c, slug, asset_id, kind, source_id=None, **fields):
    source_id = source_id or uuid.uuid4().hex
    stamp = store.now()
    c.execute('''INSERT INTO asset_sources
        (id,project_slug,asset_id,kind,label,url,file_name,file_path,mime,size_bytes,
         extracted_text,extract_status,extract_note,fetched_at,created_at,updated_at)
        VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',
        (source_id, slug, asset_id, kind, fields.get('label', ''), fields.get('url', ''),
         fields.get('file_name', ''), fields.get('file_path', ''), fields.get('mime', ''),
         int(fields.get('size_bytes', 0)), fields.get('text', ''), fields.get('status', 'empty'),
         fields.get('note', ''), fields.get('fetched_at') or None, stamp, stamp))
    store.record_event(c, slug, 'material_added',
                       f"新增素材（{kind}）：{fields.get('label') or fields.get('file_name') or fields.get('url')}")
    return source_id


def add_file(slug, asset_id, filename, data, label=''):
    if len(data) > MAX_UPLOAD_BYTES:
        raise MaterialError(f'文件超过 {MAX_UPLOAD_BYTES // 1024 // 1024}MB 上限')
    if not data:
        raise MaterialError('文件内容为空')
    name = _safe_label(filename, 'material')
    text, status, note = extract_bytes(name, data)
    with store.connection() as c:
        c.execute('BEGIN IMMEDIATE')
        _asset_exists(c, slug, asset_id)
        source_id = uuid.uuid4().hex
        relative = _write_file(slug, asset_id, source_id, name, data)
        _insert(c, slug, asset_id, 'file', source_id=source_id, label=str(label or '').strip(),
                file_name=name, file_path=relative, size_bytes=len(data),
                text=text, status=status, note=note)
        store.record_event(c, slug, 'material_added', f'上传素材：{name}（{status}）')
    if status == 'ocr_pending':
        queue_ocr(slug, asset_id, source_id)
    return {'id': source_id, 'status': status, 'note': note, 'chars': len(text)}


def add_url(slug, asset_id, url, label=''):
    target = str(url or '').strip()
    if not target:
        raise MaterialError('请填写网址')
    result = fetch_url(target)
    keep_file = result['status'] == 'ocr_pending'
    file_name = ''
    if keep_file:
        guessed = Path(urlparse(result['final_url']).path).name
        file_name = _safe_label(guessed or 'page.pdf', 'page.pdf')
        if not _safe_suffix(file_name):
            file_name += '.pdf' if result['suffix'] == '.pdf' or 'pdf' in (result['content_type'] or '') else '.img'
    with store.connection() as c:
        c.execute('BEGIN IMMEDIATE')
        _asset_exists(c, slug, asset_id)
        source_id = uuid.uuid4().hex
        relative = _write_file(slug, asset_id, source_id, file_name, result['raw']) if keep_file else ''
        _insert(c, slug, asset_id, 'url', source_id=source_id, label=str(label or '').strip(),
                url=target, file_name=file_name, file_path=relative,
                mime=result['content_type'], size_bytes=len(result['raw']) if keep_file else 0,
                text=result['text'], status=result['status'], note=result['note'],
                fetched_at=store.now())
    if result['status'] == 'ocr_pending':
        queue_ocr(slug, asset_id, source_id)
    return {'id': source_id, 'status': result['status'], 'note': result['note'],
            'chars': len(result['text']), 'final_url': result['final_url']}


def add_note(slug, asset_id, text, label=''):
    body = _normalize(text)
    if not body:
        raise MaterialError('粘贴的内容为空')
    body, note = _truncate(body)
    with store.connection() as c:
        c.execute('BEGIN IMMEDIATE')
        _asset_exists(c, slug, asset_id)
        source_id = _insert(c, slug, asset_id, 'note', label=str(label or '').strip(),
                            text=body, status='ok', note=note)
    return {'id': source_id, 'status': 'ok', 'note': note, 'chars': len(body)}


def _row(slug, asset_id, source_id):
    with store.connection() as c:
        row = c.execute('SELECT * FROM asset_sources WHERE id=? AND asset_id=? AND project_slug=?',
                        (source_id, asset_id, slug)).fetchone()
    if not row:
        raise KeyError('素材不存在')
    return dict(row)


def run_ocr(slug, asset_id, source_id):
    """Synchronous OCR pass. Used by the background worker and by tests."""
    row = _row(slug, asset_id, source_id)
    if row['kind'] == 'note':
        return {'id': source_id, 'status': 'ok', 'note': row['extract_note'],
                'chars': len(row['extracted_text'] or '')}
    path = (store.DATA / row['file_path']).resolve() if row['file_path'] else None
    if not path or store.DATA.resolve() not in path.parents or not path.is_file():
        text, status, note = '', 'failed', '原件已不在磁盘上，无法 OCR，请重新上传'
    else:
        try:
            text, status, note = ocr_bytes(row['file_name'], path.read_bytes())
        except MaterialError as exc:
            text, status, note = '', 'failed', str(exc)
    with store.connection() as c:
        c.execute('BEGIN IMMEDIATE')
        c.execute('''UPDATE asset_sources SET extracted_text=?,extract_status=?,extract_note=?,
                     updated_at=? WHERE id=?''', (text, status, note, store.now(), source_id))
        store.record_event(c, slug, 'material_ocr', f'OCR 完成（{status}，{len(text)} 字）')
    return {'id': source_id, 'status': status, 'note': note, 'chars': len(text)}


def _worker(slug, asset_id, source_id):
    try:
        run_ocr(slug, asset_id, source_id)
    except Exception:
        # 后台线程不能把异常抛给服务器；状态留在数据库里由用户重试。
        try:
            with store.connection() as c:
                c.execute('BEGIN IMMEDIATE')
                c.execute("UPDATE asset_sources SET extract_status='failed',extract_note=?,updated_at=? WHERE id=?",
                          ('OCR 进程异常结束，请点「重新提取」重试', store.now(), source_id))
        except Exception:
            pass
    finally:
        with _OCR_JOBS_LOCK:
            _OCR_JOBS.discard(source_id)

def queue_ocr(slug, asset_id, source_id):
    """Start OCR in the background, at most one job per material."""
    with _OCR_JOBS_LOCK:
        if source_id in _OCR_JOBS:
            return False
        _OCR_JOBS.add(source_id)
    thread = threading.Thread(target=_worker, args=(slug, asset_id, source_id),
                             name=f'ocr-{source_id[:8]}', daemon=True)
    with _OCR_JOBS_LOCK:
        _OCR_THREADS[source_id] = thread
    thread.start()
    return True


def wait_for_ocr(timeout=600):
    """Wait for in-flight OCR jobs.

    Tests must call this before tearing down their temporary database: a lingering worker
    would otherwise write into whatever ``store.DB`` points at next.
    """
    deadline = time.time() + timeout
    while True:
        with _OCR_JOBS_LOCK:
            threads = list(_OCR_THREADS.values())
        alive = [t for t in threads if t.is_alive()]
        if not alive:
            with _OCR_JOBS_LOCK:
                _OCR_THREADS.clear()
            return True
        if time.time() > deadline:
            return False
        for thread in alive:
            thread.join(timeout=0.5)


def reset_stuck_ocr():
    """A restarted service cannot have OCR threads still running."""
    with store.connection() as c:
        c.execute('BEGIN IMMEDIATE')
        changed = c.execute(
            "UPDATE asset_sources SET extract_status='pending',"
            "extract_note='服务已重启，OCR 未完成，请点「重新提取」',updated_at=? "
            "WHERE extract_status='ocr_pending'", (store.now(),)).rowcount
    return changed


def refresh(slug, asset_id, source_id):
    row = _row(slug, asset_id, source_id)
    if row['kind'] == 'url':
        result = fetch_url(row['url'])
        text, status, note = result['text'], result['status'], result['note']
        mime, fetched = result['content_type'], store.now()
        file_path = row['file_path']
        if status == 'ocr_pending':
            file_path = _write_file(slug, asset_id, source_id, row['file_name'] or result['suffix'] or 'page',
                                    result['raw'])
    elif row['kind'] == 'file':
        path = (store.DATA / row['file_path']).resolve()
        if store.DATA.resolve() not in path.parents or not path.is_file():
            raise MaterialError('原件已不在磁盘上，请重新上传')
        text, status, note = extract_bytes(row['file_name'], path.read_bytes())
        mime, fetched, file_path = row['mime'], row['fetched_at'], row['file_path']
    else:
        text, status, note = row['extracted_text'], 'ok', row['extract_note']
        mime, fetched, file_path = row['mime'], row['fetched_at'], row['file_path']
    with store.connection() as c:
        c.execute('BEGIN IMMEDIATE')
        _asset_exists(c, slug, asset_id)
        c.execute('''UPDATE asset_sources SET extracted_text=?,extract_status=?,extract_note=?,
                     mime=?,fetched_at=?,file_path=?,updated_at=? WHERE id=?''',
                  (text, status, note, mime, fetched, file_path, store.now(), source_id))
        store.record_event(c, slug, 'material_refreshed', f'重新提取素材（{status}）')
    if status == 'ocr_pending':
        queue_ocr(slug, asset_id, source_id)
    return {'id': source_id, 'status': status, 'note': note, 'chars': len(text)}


def delete(slug, asset_id, source_id):
    with store.connection() as c:
        c.execute('BEGIN IMMEDIATE')
        row = c.execute('SELECT file_path FROM asset_sources WHERE id=? AND asset_id=? AND project_slug=?',
                        (source_id, asset_id, slug)).fetchone()
        if not row:
            raise KeyError('素材不存在')
        c.execute('DELETE FROM asset_sources WHERE id=?', (source_id,))
        store.record_event(c, slug, 'material_deleted', '删除素材')
    if row['file_path']:
        path = (store.DATA / row['file_path']).resolve()
        if store.DATA.resolve() in path.parents:
            try:
                path.unlink()
            except OSError:
                pass
    return {'message': '已删除素材'}


def delete_files_for(asset_ids):
    """Remove material files for the given assets (called before the rows cascade away)."""
    if not asset_ids:
        return 0
    removed = 0
    for asset_id in asset_ids:
        directory = None
        with store.connection() as c:
            rows = [dict(r) for r in c.execute(
                'SELECT file_path FROM asset_sources WHERE asset_id=?', (asset_id,))]
        for row in rows:
            if not row['file_path']:
                continue
            path = (store.DATA / row['file_path']).resolve()
            if store.DATA.resolve() not in path.parents:
                continue
            directory = path.parent
            try:
                path.unlink()
                removed += 1
            except OSError:
                pass
        if directory and directory.is_dir():
            try:
                directory.rmdir()
            except OSError:
                pass
    return removed


def file_path_for(slug, asset_id, source_id):
    with store.connection() as c:
        row = c.execute('''SELECT file_path,file_name FROM asset_sources
                           WHERE id=? AND asset_id=? AND project_slug=?''',
                        (source_id, asset_id, slug)).fetchone()
    if not row or not row['file_path']:
        raise KeyError('该素材没有原件可下载')
    path = (store.DATA / row['file_path']).resolve()
    if store.DATA.resolve() not in path.parents or not path.is_file():
        raise KeyError('原件已不在磁盘上')
    return path, row['file_name'] or path.name


# --------------------------------------------------------------------------- 读取与摘要

def _public(row, excerpt_chars=0):
    text = row['extracted_text'] or ''
    item = {'id': row['id'], 'kind': row['kind'], 'label': row['label'],
            'url': row['url'], 'file_name': row['file_name'],
            'mime': row['mime'], 'size_bytes': row['size_bytes'],
            'status': row['extract_status'], 'note': row['extract_note'],
            'chars': len(text), 'fetched_at': row['fetched_at'],
            'has_file': bool(row['file_path']),
            'created_at': row['created_at'], 'updated_at': row['updated_at']}
    if excerpt_chars:
        item['excerpt'] = text[:excerpt_chars]
    return item


def list_sources(slug, asset_id, excerpt_chars=0):
    with store.connection() as c:
        rows = [dict(r) for r in c.execute(
            'SELECT * FROM asset_sources WHERE project_slug=? AND asset_id=? ORDER BY created_at',
            (slug, asset_id))]
    return [_public(row, excerpt_chars) for row in rows]


def digest(slug, asset_id, per_source_chars=3000, total_chars=18000):
    """Usable material excerpts for drafting. Returns (items, combined_text)."""
    items, combined, used = [], [], 0
    for row in list_sources(slug, asset_id, excerpt_chars=per_source_chars):
        if row['status'] != 'ok' or row['chars'] <= 0:
            continue
        head = row['label'] or row['file_name'] or row['url']
        block = f'【{head}】\n{row["excerpt"]}'
        if used + len(block) > total_chars:
            break
        used += len(block)
        item = {key: row[key] for key in
                ('id', 'kind', 'label', 'url', 'file_name', 'status', 'chars')}
        item['excerpt'] = row['excerpt']
        items.append(item)
        combined.append(block)
    return items, '\n\n'.join(combined)


def pending_count(slug, asset_id):
    return sum(1 for row in list_sources(slug, asset_id)
               if row['status'] in {'ocr_pending', 'pending'})


def source_lines(sources):
    """Human-reviewable source lines for the facts field."""
    lines = []
    for row in sources:
        if row['status'] != 'ok':
            continue
        if row['kind'] == 'url':
            lines.append(f"来源：{row['url']}（已抓取 {row['chars']} 字，待人工确认）")
        elif row['kind'] == 'file':
            lines.append(f"文件：{row['file_name']}（已提取 {row['chars']} 字，待人工确认）")
        else:
            lines.append(f"来源：粘贴素材「{row['label'] or '未命名'}」（已录入 {row['chars']} 字，待人工确认）")
    return lines
